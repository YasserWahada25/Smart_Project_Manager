from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

NAVY = RGBColor(0x0F, 0x1B, 0x3D)
TEAL = RGBColor(0x14, 0xB8, 0xA6)
WHITE = RGBColor(255, 255, 255)
LIGHT = RGBColor(0xF3, 0xF6, 0xFA)
GREY = RGBColor(0x4B, 0x55, 0x63)
p = Presentation()
p.slide_width = Inches(13.333)
p.slide_height = Inches(7.5)
blank = p.slide_layouts[6]


def rect(s, x, y, w, h, c, shape=MSO_SHAPE.RECTANGLE):
    r = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = c
    r.line.fill.background()
    return r


def text(s, x, y, w, h, t, size=18, color=GREY, bold=False, align=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    lines = t if isinstance(t, list) else [t]
    for i, l in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = l
        para.font.size = Pt(size)
        para.font.color.rgb = color
        para.font.bold = bold
        para.space_after = Pt(8)
        if align:
            para.alignment = align
    return tb


def slide(title, n):
    s = p.slides.add_slide(blank)
    rect(s, 0, 0, 13.333, 7.5, LIGHT)
    rect(s, 0, 0, 13.333, 1.25, NAVY)
    rect(s, 0, 1.25, 13.333, 0.07, TEAL)
    text(s, 0.6, 0.3, 11.5, 0.8, title, 32, WHITE, True)
    text(s, 12.2, 6.95, 1, 0.4, str(n), 12, GREY)
    return s


def bullets(s, items, x=0.7, y=1.8, w=11.9, size=20):
    for i, (h, d) in enumerate(items):
        yy = y + i * 0.85
        rect(s, x, yy + 0.1, 0.14, 0.5, TEAL)
        tb = s.shapes.add_textbox(Inches(x + 0.3), Inches(yy), Inches(w), Inches(0.8))
        tf = tb.text_frame
        tf.word_wrap = True
        para = tf.paragraphs[0]
        r = para.add_run()
        r.text = h + "  "
        r.font.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = NAVY
        r = para.add_run()
        r.text = d
        r.font.size = Pt(size - 2)
        r.font.color.rgb = GREY


def table(s, rows, x, y, w, colw=None, size=16, rh=0.5, total=True):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(rh * len(rows))).table
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            c = t.cell(i, j)
            c.text = v
            f = c.text_frame.paragraphs[0].font
            f.size = Pt(size)
            c.fill.solid()
            if i == 0:
                c.fill.fore_color.rgb = NAVY
                f.color.rgb = WHITE
                f.bold = True
            elif total and i == len(rows) - 1:
                c.fill.fore_color.rgb = RGBColor(0xD1, 0xFA, 0xE5)
                f.bold = True
                f.color.rgb = NAVY
            else:
                c.fill.fore_color.rgb = WHITE
                f.color.rgb = GREY
    if colw:
        for j, cw in enumerate(colw):
            t.columns[j].width = Inches(cw)


def img(s, f, x, y, maxw, maxh):
    w, h = Image.open(f).size
    r = min(maxw / w, maxh / h)
    s.shapes.add_picture(f, Inches(x), Inches(y), Inches(w * r), Inches(h * r))


def note(s, t):
    text(s, 0.7, 6.7, 11, 0.5, t, 14, GREY)


# 1 titre
s = p.slides.add_slide(blank)
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 0.8, 2.2, 0.12, 2.4, TEAL)
text(s, 1.2, 1.3, 11, 0.5, "TRAVAIL DE RECHERCHE  •  TEST ET VALIDATION", 16, TEAL, True)
text(s, 1.2, 2.1, 11, 1.5, "Outils de test statique", 54, WHITE, True)
text(s, 1.2, 3.4, 11, 1, "Application au projet PI – FakarniApp", 30, WHITE)
text(s, 1.2, 4.2, 11, 0.6, "SonarCloud / SonarQube + GitHub Actions + JUnit + JaCoCo", 20, TEAL)
text(s, 1.2, 6.0, 11, 0.8, "Classe 4SAE5  •  Esprit_PI_4SAE5_2026_FakarniApp  •  2026-2027", 16, RGBColor(0xCB, 0xD5, 0xE1))

# 2 contexte
s = slide("Contexte et objectif", 2)
bullets(s, [
    ("Tests statiques :", "analyse du code source sans l'exécuter."),
    ("Ce qu'ils détectent :", "défauts de qualité, risques de sécurité, duplications, maintenabilité."),
    ("Dans FakarniApp :", "analyse intégrée à la CI/CD via SonarCloud."),
    ("Tests dynamiques :", "exécutés avec JUnit et suivis par JaCoCo."),
    ("Objectif :", "présenter un outil statique et montrer son intégration réelle dans le projet PI."),
])

# 3 sonar
s = slide("SonarCloud / SonarQube", 3)
text(s, 0.7, 1.6, 12, 1, "SonarQube est une plateforme d'analyse statique ; SonarCloud est sa version hébergée dans le cloud. Ils détectent automatiquement les problèmes de fiabilité, sécurité et maintenabilité.", 18)
items = [("Bugs", "Comportement incorrect"), ("Vulnerabilities", "Failles de sécurité"), ("Code Smells", "Code difficile à maintenir"),
         ("Security Hotspots", "Zones à vérifier à la main"), ("Duplications", "Code répété"), ("Coverage", "Importée depuis JaCoCo")]
for i, (h, d) in enumerate(items):
    x = 0.7 + (i % 3) * 4.1
    y = 3.2 + (i // 3) * 1.8
    rect(s, x, y, 3.8, 1.5, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, y, 0.12, 1.5, TEAL)
    text(s, x + 0.3, y + 0.2, 3.4, 0.5, h, 22, NAVY, True)
    text(s, x + 0.3, y + 0.8, 3.4, 0.6, d, 16)

# 4 pourquoi
s = slide("Pourquoi SonarCloud pour FakarniApp ?", 4)
bullets(s, [
    ("Architecture :", "plusieurs microservices Java / Spring Boot."),
    ("Intégration :", "directe avec GitHub Actions."),
    ("Couverture :", "exploite les rapports JaCoCo."),
    ("Démarche :", "parfaitement adaptée à notre CI/CD."),
], y=2.0)

# 5 architecture
s = slide("Architecture de validation", 5)
table(s, [["Outil / composant", "Rôle dans le projet"],
          ["GitHub Actions", "Automatise build, tests, analyse statique et Docker"],
          ["JUnit / Spring Boot Test", "Tests dynamiques des contrôleurs, services et contexte"],
          ["Maven Surefire", "Génère les rapports XML de tests"],
          ["JaCoCo", "Mesure la couverture de code Java"],
          ["SonarCloud", "Analyse statique qualité et sécurité"],
          ["Docker", "Construit et publie l'image après validation"]],
      0.7, 1.7, 11.9, [3.8, 8.1], 17, 0.62, False)

# 6 comparaison
s = slide("Tests dynamiques vs analyse statique", 6)
cols = [("Tests dynamiques", ["JUnit / Spring Boot Test", "Exécutent le code", "Vérifient comportements et composants", "JaCoCo mesure la couverture"]),
        ("Analyse statique", ["SonarCloud / SonarQube", "Analyse sans exécuter les scénarios", "Détecte bugs, vulnérabilités, code smells, duplications", "Exploite la couverture JaCoCo"])]
for k, (h, its) in enumerate(cols):
    x = 0.7 + k * 6.1
    rect(s, x, 1.7, 5.8, 4.8, WHITE, MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, 1.7, 5.8, 0.8, TEAL if k else NAVY)
    text(s, x + 0.3, 1.82, 5.2, 0.6, h, 24, WHITE, True)
    text(s, x + 0.3, 2.8, 5.2, 3.5, ["• " + i for i in its], 18)
note(s, "Les deux approches sont complémentaires.")

# 7 pipeline
s = slide("Application réelle : pipeline GitHub Actions", 7)
img(s, "image1.png", 0.5, 1.5, 7.6, 5.2)
steps = ["Build", "Test + JaCoCo", "Publish test results", "SonarCloud analysis", "Package JAR", "Docker Build & Push"]
for i, t in enumerate(steps):
    y = 1.6 + i * 0.82
    rect(s, 8.4, y, 4.4, 0.68, TEAL if "Sonar" in t else NAVY, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, 8.55, y + 0.1, 4.1, 0.5, f"{i + 1}. {t}", 17, WHITE, True)
note(s, "Service Suivi Engagement – job terminé avec succès, toutes les étapes validées.")


def res(n, title, file, rows, cap, colw):
    s = slide(title, n)
    img(s, file, 0.5, 1.6, 7.2, 4.9)
    table(s, rows, 8.0, 1.8, 4.9, colw, 15, 0.5)
    note(s, cap)


res(8, "Suivi Engagement : 3 tests réussis", "image2.png",
    [["Suite", "OK", "KO", "Ign."], ["PatientEngagementControllerTest", "1", "0", "0"], ["PatientEngagementServiceTest", "1", "0", "0"],
     ["SuiviEngagement…Tests", "1", "0", "0"], ["TOTAL", "3", "0", "0"]],
    "Tests de base (chargement de classes, contexte Spring) : utiles mais ne remplacent pas des tests métier détaillés.", [2.8, 0.7, 0.7, 0.7])

s = slide("User Service : 15 tests réussis", 9)
img(s, "image3.jpg", 0.5, 1.6, 7.2, 4.9)
text(s, 8.0, 1.7, 4.9, 5, ["• AuthControllerTest : 2 (login, révocation refresh token)", "• InternalUserControllerTest : 1", "• UserControllerTest : 3",
                           "• JwtServiceTest, SessionCookieServiceTest : sécurité et sessions", "• AuthServiceTest, EmailServiceTest, UserServiceTest : couche service", "• UserApplicationTests : contexte"], 16)
note(s, "Scénarios précis sur l'authentification, au-delà des simples tests de chargement.")

res(10, "Group Service : 8 tests réussis", "image4.jpg",
    [["Composant", "OK"], ["JwtServiceTest", "1"], ["FileUploadControllerTest", "1"], ["GroupControllerTest", "3"],
     ["GroupApplicationTests", "1"], ["GroupServiceTest", "2"], ["TOTAL", "8"]],
    "Configuration JWT, upload de fichiers, contrôleur, service et contexte Spring Boot.", [3.7, 1.2])
res(11, "Event Service : 10 tests réussis", "image5.jpg",
    [["Suite", "OK"], ["EmailControllerTest", "1"], ["EventControllerTest", "3"], ["MapsControllerTest", "1"],
     ["EventServiceApplicationTests", "1"], ["EmailServiceTest", "1"], ["EventServiceTest", "3"], ["TOTAL", "10"]],
    "Contrôleurs et services liés aux événements, à l'email et aux cartes.", [3.7, 1.2])

# 12 bilan
s = slide("Bilan global : 36 tests réussis, 0 échec", 12)
table(s, [["Microservice", "Réussis", "Échecs", "Ignorés"], ["Suivi Engagement", "3", "0", "0"], ["User Service", "15", "0", "0"],
          ["Group Service", "8", "0", "0"], ["Event Service", "10", "0", "0"], ["TOTAL", "36", "0", "0"]],
      0.7, 1.8, 7.4, [3.4, 1.3, 1.4, 1.3], 18, 0.7)
rect(s, 8.6, 1.8, 4.1, 4.2, NAVY, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 8.6, 2.1, 4.1, 1.5, "36", 90, TEAL, True, 2)
text(s, 8.8, 4.2, 3.7, 1.5, "tests dynamiques réussis sur 4 microservices", 18, WHITE, False, 2)

# 13 interprétation
s = slide("Interprétation par rapport au test statique", 13)
bullets(s, [
    ("36 tests visibles :", "dynamiques (exécution du code, rapports Surefire)."),
    ("Élément statique :", "l'étape « SonarCloud analysis » du pipeline."),
    ("À compléter :", "capture du tableau de bord SonarCloud (Bugs, Vulnerabilities, Code Smells, Hotspots, Coverage, Duplications, Quality Gate)."),
    ("Honnêteté :", "ces métriques ne figurent pas dans les captures, elles ne sont donc pas inventées."),
], y=1.9)

# 14 conclusion
s = slide("Conclusion", 14)
bullets(s, [
    ("SonarCloud", "est notre outil statique, équivalent cloud de SonarQube."),
    ("Intégré", "dans GitHub Actions après les tests et JaCoCo."),
    ("36 tests", "dynamiques réussis sur quatre microservices."),
    ("Complémentarité :", "JUnit + JaCoCo vérifient le comportement, SonarCloud la qualité interne."),
    ("Résultat :", "fiabilité, maintenabilité et sécurité accrues, puis packaging et build Docker."),
])

# 15 merci
s = p.slides.add_slide(blank)
rect(s, 0, 0, 13.333, 7.5, NAVY)
rect(s, 5.4, 3.9, 2.5, 0.08, TEAL)
text(s, 0, 2.6, 13.333, 1.2, "Merci pour votre attention", 48, WHITE, True, 2)
text(s, 0, 4.2, 13.333, 0.8, "Des questions ?", 26, TEAL, False, 2)
p.save("FakarniApp_Tests_Statiques.pptx")
