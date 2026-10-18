# -*- coding: utf-8 -*-
"""Livrable Word de l'Atelier 2 (python-docx) : packages, black / white box, interfaces, méthodes."""
import os
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from spec import PACKAGES, DEPENDENCIES, INTERFACES

HERE = os.path.dirname(os.path.abspath(__file__))
IMG = os.path.join(HERE, "img")
OUT = os.path.join(HERE, "..", "Atelier2_Packages_Composants_Classes.docx")
FONT = "Calibri"

CLASS_ROLE = {
    "AuthService": "Inscription et connexion ; construit la réponse (jeton + utilisateur).",
    "TokenService": "Signe et vérifie le jeton JWT (HS256).",
    "AccessGuard": "Vérifie le jeton, charge l'utilisateur et contrôle son rôle.",
    "UserService": "Administration des comptes : liste, statut actif, rôle ; annuaire des développeurs.",
    "ProfileService": "Profil de l'utilisateur connecté : informations, mot de passe, compétences.",
    "User": "Compte utilisateur (rôle ADMIN, PROJECT_MANAGER ou DEVELOPER).",
    "Skill": "Compétence d'un développeur (nom, niveau, années d'expérience).",
    "ProjectService": "Cycle de vie des projets et gestion des membres de l'équipe.",
    "Project": "Projet logiciel : dates, statut, technologies, responsable, membres.",
    "ProjectAccessService": "Règles d'accès au projet : visibilité, responsable, projet archivé.",
    "EventBus": "Publie les événements métier et notifie les abonnés (patron observateur).",
    "DomainEvent": "Événement métier : type, acteur, objets concernés, détails.",
    "SprintService": "Création, modification, changement de statut et suppression des sprints ; statistiques.",
    "Sprint": "Itération d'un projet (PLANNED, ACTIVE, COMPLETED, CANCELLED).",
    "SprintStats": "Mesures d'avancement d'un sprint en points d'effort.",
    "TaskService": "Création, modification, workflow, affectation et consultation des tâches.",
    "Task": "Tâche de développement : priorité, complexité, statut, assigné, sprint.",
    "TaskWorkflow": "Table des transitions de statut autorisées (TODO → … → DONE, BLOCKED).",
    "BoardService": "Construit le tableau Kanban d'un projet.",
    "Board": "Tâches regroupées par colonne de statut.",
    "CommentService": "Commentaires des tâches : lecture, ajout, modification, modération.",
    "Comment": "Commentaire écrit par un membre sur une tâche.",
    "ActivityService": "Enregistre les événements dans l'historique et le consulte.",
    "Activity": "Entrée de l'historique d'activité du projet.",
    "NotificationService": "Transforme les activités en notifications et gère la boîte de réception.",
    "Notification": "Message adressé à un utilisateur (lu / non lu).",
    "DashboardService": "Calcule les indicateurs globaux et par projet.",
    "Indicators": "Ensemble d'indicateurs affichés sur un tableau de bord.",
    "SearchService": "Recherche de projets et de tâches visibles par l'utilisateur.",
    "SearchResult": "Résultat d'une recherche (projets et tâches).",
    "AiClientService": "Client HTTP du service IA : jeton de service, délai limité, erreurs mappées.",
    "AiPlanService": "AI-01 : extrait le document, demande le plan, valide la réponse, crée sprints et tâches.",
    "ProjectPlan": "Plan proposé par l'IA : sprints, backlog, avertissements.",
    "AiRecommendationService": "AI-02 : classe les membres selon compétences, charge et expérience.",
    "DeveloperRecommendation": "Score et explication pour un développeur candidat.",
    "AiRiskService": "AI-03 : prédit le risque de retard d'un sprint actif.",
    "SprintRisk": "Niveau de risque, probabilité et facteurs d'un sprint.",
    "AiAssistantService": "AI-04 : dialogue avec l'assistant, outils de lecture, actions proposées puis confirmées.",
    "ProposedAction": "Modification proposée par l'assistant, appliquée seulement après confirmation.",
}

SUB_ROLE = {
    ("Identity", "Authentication"): "Inscription, connexion, jeton JWT.",
    ("Identity", "Authorization"): "Authentification de la requête et contrôle des rôles.",
    ("Identity", "Users"): "Administration des comptes, profil et compétences.",
    ("Identity", "Domain model"): "Entités User et Skill.",
    ("Projects", "Project management"): "CRUD des projets, statuts, membres de l'équipe.",
    ("Projects", "Access control"): "Visibilité et droits sur un projet (réutilisés par tous).",
    ("Projects", "Domain events"): "Publication et abonnement aux événements métier.",
    ("Sprints", "Lifecycle"): "CRUD et cycle de vie des sprints.",
    ("Sprints", "Statistics"): "Avancement d'un sprint.",
    ("Tasks", "Task management"): "CRUD, affectation et statuts des tâches.",
    ("Tasks", "Workflow"): "Règles de transition entre statuts.",
    ("Tasks", "Board"): "Vue Kanban par colonne.",
    ("Collaboration", "Comments"): "Commentaires sur les tâches.",
    ("Collaboration", "History"): "Journal des événements du projet.",
    ("Collaboration", "Notifications"): "Notifications générées à partir des activités.",
    ("Insight", "Dashboard"): "Indicateurs et graphiques de pilotage.",
    ("Insight", "Search"): "Recherche globale.",
    ("AI", "Gateway"): "Passerelle vers le service IA externe.",
    ("AI", "Planning (AI-01)"): "Génération du plan sprints / tâches.",
    ("AI", "Recommendation (AI-02)"): "Recommandation du meilleur développeur.",
    ("AI", "Risk (AI-03)"): "Prédiction du risque de retard.",
    ("AI", "Assistant (AI-04)"): "Assistant conversationnel de gestion.",
}

FR_TABLE = [
    ("FR-01", "Authentification", "Inscription, connexion par JWT, routes protégées, autorisation par rôle."),
    ("FR-02", "Gestion des utilisateurs", "L'administrateur gère les comptes, les rôles et l'activation."),
    ("FR-03", "Profils utilisateurs", "Consulter et modifier son profil, changer son mot de passe."),
    ("FR-04", "Compétences des développeurs", "Gérer nom, niveau et années d'expérience (utilisées par AI-02)."),
    ("FR-05", "Gestion des projets", "CRUD des projets, cycle de vie des statuts, membres."),
    ("FR-06", "Gestion des équipes", "Ajouter et retirer des membres, annuaire des développeurs."),
    ("FR-07", "Gestion des sprints", "CRUD des sprints d'un projet et cycle de vie."),
    ("FR-08", "Gestion des tâches", "CRUD, affectation, transitions du workflow."),
    ("FR-09", "Tableau Kanban", "Visualiser les tâches par colonne, tâches BLOCKED identifiables."),
    ("FR-10", "Commentaires", "Commenter les tâches."),
    ("FR-11", "Historique d'activité", "Journaliser les événements de tâches, sprints et projets."),
    ("FR-12", "Notifications", "Informer les utilisateurs des événements qui les concernent."),
    ("FR-13", "Tableau de bord", "Indicateurs et graphiques par rôle et par projet."),
    ("FR-14", "Recherche et filtres", "Rechercher et filtrer projets, tâches, développeurs."),
    ("FR-15", "Fonctionnalités d'IA", "AI-01 plan, AI-02 recommandation, AI-03 risque de retard, AI-04 assistant."),
]

TRACE = [
    ("FR-01", "Identity", "Authentication", "AuthService, TokenService, AccessGuard"),
    ("FR-02", "Identity", "Users", "UserService, User"),
    ("FR-03", "Identity", "Users", "ProfileService, User"),
    ("FR-04", "Identity", "Users, Domain model", "ProfileService, Skill"),
    ("FR-05", "Projects", "Project management, Access control", "ProjectService, Project, ProjectAccessService"),
    ("FR-06", "Projects", "Project management", "ProjectService (addMember, removeMember), UserService"),
    ("FR-07", "Sprints", "Lifecycle, Statistics", "SprintService, Sprint, SprintStats"),
    ("FR-08", "Tasks", "Task management, Workflow", "TaskService, Task, TaskWorkflow"),
    ("FR-09", "Tasks", "Board", "BoardService, Board"),
    ("FR-10", "Collaboration", "Comments", "CommentService, Comment"),
    ("FR-11", "Collaboration", "History", "ActivityService, Activity, EventBus"),
    ("FR-12", "Collaboration", "Notifications", "NotificationService, Notification"),
    ("FR-13", "Insight", "Dashboard", "DashboardService, Indicators"),
    ("FR-14", "Insight", "Search", "SearchService, SearchResult"),
    ("FR-15", "AI", "Gateway, Planning, Recommendation, Risk, Assistant",
     "AiClientService, AiPlanService, AiRecommendationService, AiRiskService, AiAssistantService"),
]

doc = Document()
fig_no = [0]
tab_no = [0]


# ------------------------------------------------------------------ styles
def set_font(style, size, bold=False, italic=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), FONT)


st = doc.styles
set_font(st["Normal"], 11)
st["Normal"].paragraph_format.space_after = Pt(6)
st["Normal"].paragraph_format.line_spacing = 1.1
for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
    set_font(st[name], size, bold=True)
    st[name].paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
    st[name].paragraph_format.space_after = Pt(6)
    st[name].paragraph_format.keep_with_next = True
set_font(st["List Bullet"], 11)
set_font(st["Caption"], 10, italic=True)


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def para(text="", bold=False, italic=False, align=None, size=None, after=None, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def rich(parts, style=None):
    """parts : liste de (texte, gras)."""
    p = doc.add_paragraph(style=style)
    for t, b in parts:
        r = p.add_run(t)
        r.bold = b
    return p


def bullets(items):
    for it in items:
        if isinstance(it, tuple):
            rich([(it[0], True), (it[1], False)], style="List Bullet")
        else:
            doc.add_paragraph(it, style="List Bullet")


def h(text, level=1, brk=False):
    p = doc.add_heading(text, level=level)
    if brk:
        p.paragraph_format.page_break_before = True
    return p


def caption(text, kind="Figure"):
    n = fig_no if kind == "Figure" else tab_no
    n[0] += 1
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("%s %d – %s" % (kind, n[0], text))
    r.italic = True
    if kind == "Tableau":
        p.paragraph_format.keep_with_next = True
    return p


def table(headers, rows, widths, font_size=10, mono_cols=()):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, hd in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = Cm(widths[i])
        shade(c, "E7E6E6")
        c.paragraphs[0].text = ""
        r = c.paragraphs[0].add_run(hd)
        r.bold = True
        r.font.size = Pt(font_size)
    trPr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trPr.append(th)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = Cm(widths[i])
            lines = val if isinstance(val, list) else [val]
            for k, ln in enumerate(lines):
                p = cells[i].paragraphs[0] if k == 0 else cells[i].add_paragraph()
                p.paragraph_format.space_after = Pt(1)
                r = p.add_run(ln)
                r.font.size = Pt(font_size)
                if i in mono_cols:
                    r.font.name = "Consolas"
                    r.font.size = Pt(font_size - 1)
        tr = t.rows[-1]._tr
        trPr = tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        cs.set(qn("w:val"), "true")
        trPr.append(cs)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(path, width_cm, cap, max_h_cm=None):
    from PIL import Image
    w, hh = Image.open(path).size
    if max_h_cm and width_cm * hh / w > max_h_cm:
        width_cm = max_h_cm * w / hh
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width_cm))
    caption(cap)


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_field(run, instr):
    for typ, text in (("begin", None), (None, instr), ("end", None)):
        if typ:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), typ)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def footer(section):
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Atelier 2 – Smart Project Manager – page ")
    r.font.size = Pt(9)
    r2 = p.add_run()
    r2.font.size = Pt(9)
    add_field(r2, "PAGE")


def new_section(landscape):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if landscape:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(1.6)
        s.top_margin = s.bottom_margin = Cm(1.5)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.2)
        s.top_margin = s.bottom_margin = Cm(2.2)
    footer(s)
    return s


# ------------------------------------------------------------------ chiffres
n_sub = sum(len(p["subs"]) for p in PACKAGES)
n_cls = sum(len(s["classes"]) for p in PACKAGES for s in p["subs"])
n_meth = sum(len(c.get("methods", [])) for p in PACKAGES for s in p["subs"] for c in s["classes"])

# ================================================================== page de garde
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.2)
footer(sec)
for _ in range(5):
    para()
para("ESPRIT University", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=16)
para("ATELIER 2 – ARCHITECTURE LOGICIELLE", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, after=30)
para("Smart Project Manager", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=30, after=10)
para("Modélisation des exigences fonctionnelles par packages, sous-packages, composants et classes",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=15, after=40)
para("Travail de groupe – Atelier 2 – Semaine 2", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
para("Technologies : Angular, Express.js, MongoDB, FastAPI", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, after=40)
para("Réalisé par : ……………………………………………………………", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
para("Encadrant : ……………………………………………………………", align=WD_ALIGN_PARAGRAPH.CENTER, size=12)
page_break()

h("Plan du document")
for t in ("1. Objectifs de l'atelier et conventions de notation",
          "2. Rappel des exigences fonctionnelles (Atelier 1)",
          "3. Étape 1 – Modélisation par packages",
          "4. Étape 2 – Raffinement : black box et white box",
          "5. Étape 3 – Interactions : dépendances et interfaces",
          "6. Étape 4 – Méthodes des classes",
          "7. Traçabilité des exigences",
          "8. Conclusion"):
    para(t, after=3)
page_break()

# ================================================================== 1
h("1. Objectifs de l'atelier et conventions de notation")
para("Cet atelier fait passer le projet Smart Project Manager des exigences fonctionnelles (Atelier 1) à une "
     "première architecture logicielle. Les objectifs sont les suivants :")
bullets([
    "Modéliser les exigences fonctionnelles de l'atelier 1 (FR-01 à FR-15) par des packages.",
    "Raffiner le diagramme obtenu avec des sous-packages, des composants et des classes, en appliquant les "
    "concepts de black box et de white box.",
    "Définir les interactions par des liens de dépendance et des interfaces.",
    "Ajouter des méthodes aux différentes classes.",
])
h("1.1 Démarche", 2)
table(["Étape", "Travail réalisé", "Livrable"], [
    ["1", "Regrouper les 15 exigences en packages cohérents.", "Figure 1, tableau 2"],
    ["2", "Décomposer chaque package : sous-packages, composants, classes (white box).", "Figures 2 à 8, tableaux 3 à 10"],
    ["3", "Relier les packages par des dépendances et des interfaces fournies / requises.", "Figure 9, tableaux 11 et 12"],
    ["4", "Ajouter les méthodes à chaque classe.", "Figures 10 à 16, tableaux 13 à 19"],
], [2, 11.5, 3.1])
h("1.2 Conventions de notation", 2)
para("Conformément à la consigne, tous les diagrammes sont simples et sans couleur (noir sur fond blanc). "
     "La notation UML utilisée est la suivante :")
table(["Élément", "Représentation"], [
    ["Package / sous-package", "Dossier : rectangle surmonté d'un onglet portant le nom."],
    ["Composant", "Rectangle avec le stéréotype «component» et le pictogramme de composant."],
    ["Classe", "Rectangle à trois compartiments : nom (avec stéréotype), attributs, méthodes. "
               "Visibilité : « - » privé, « + » public."],
    ["Interface fournie", "Trait terminé par un cercle (—○), nom de l'interface sous le cercle."],
    ["Interface requise", "Trait terminé par un demi-cercle ouvert, nom de l'interface et fournisseur entre parenthèses."],
    ["Dépendance", "Flèche en tirets à pointe ouverte (- - ->) : la source utilise la cible ; le nom de l'interface utilisée figure sur la flèche."],
    ["Système externe", "Rectangle avec le stéréotype «external system»."],
], [4.5, 12.1])

# ================================================================== 2
h("2. Rappel des exigences fonctionnelles (Atelier 1)", brk=True)
para("Smart Project Manager est une application web de gestion de projets logiciels enrichie par l'intelligence "
     "artificielle. Trois rôles l'utilisent : ADMIN (gère les utilisateurs), PROJECT_MANAGER (gère projets, sprints, "
     "tâches et utilise l'IA) et DEVELOPER (travaille sur ses tâches, commente, gère son profil et ses compétences).")
caption("Exigences fonctionnelles de l'atelier 1", kind="Tableau")
table(["ID", "Module", "Résumé"], [list(r) for r in FR_TABLE], [2, 5, 9.6])
para("Deux exigences non fonctionnelles guident le découpage : NFR-01 (code en couches, modulaire) et NFR-05 "
     "(frontend, backend et service IA indépendants et déployables séparément). Une modification d'un module ne doit pas "
     "impacter les modules indépendants.")

# ================================================================== 3
h("3. Étape 1 – Modélisation par packages")
para("Chaque package regroupe des exigences qui évoluent ensemble, manipulent les mêmes données et sont utilisées par "
     "les mêmes acteurs (forte cohésion), tout en limitant les échanges avec les autres packages (faible couplage). "
     "Les 15 exigences sont réparties en 7 packages.")
figure(os.path.join(IMG, "fig01_packages.png"), 16.6, "Diagramme de packages : un package par groupe d'exigences")
caption("Packages et exigences couvertes", kind="Tableau")
table(["Package", "Exigences", "Responsabilité"],
      [[p["name"], p["fr"], p["role"]] for p in PACKAGES], [2.8, 5.0, 8.8])
h("3.1 Critères de regroupement", 2)
bullets([
    ("Identity : ", "FR-01 à FR-04 partagent l'entité utilisateur ; tout contrôle d'accès en dépend."),
    ("Projects : ", "FR-05 et FR-06 forment l'unité de base ; l'équipe n'a de sens que dans un projet."),
    ("Sprints et Tasks : ", "séparés car le cycle de vie d'un sprint (planification) diffère de celui d'une tâche (workflow, Kanban)."),
    ("Collaboration : ", "FR-10 à FR-12 sont des services transverses qui réagissent aux autres modules."),
    ("Insight : ", "FR-13 et FR-14 ne font que lire et agréger les données des autres packages."),
    ("AI : ", "FR-15 est isolé : une panne de l'IA ne doit pas arrêter la gestion de projets."),
])

# ================================================================== 4  (paysage)
new_section(True)
h("4. Étape 2 – Raffinement : black box et white box")
h("4.1 Principe", 2)
para("Black box : un bloc est vu de l'extérieur, par ce qu'il fait et par les interfaces qu'il expose, sans connaître "
     "son contenu. White box : on ouvre le bloc pour montrer sa composition interne. Le raffinement se fait donc en "
     "quatre niveaux : package (black box) → sous-packages → composants → classes (white box).")
caption("Vue black box des packages", kind="Tableau")
table(["Package", "Ce que fait le bloc (vu de l'extérieur)", "Interfaces fournies"], [
    [p["name"], p["role"], [n for n, _ in p["provided"]]] for p in PACKAGES
], [3.2, 15.0, 7.2], font_size=10)
h("4.2 Décomposition de chaque package (white box)", 2)
para("Pour chaque package, le diagramme montre ses sous-packages, le composant qui réalise chacun d'eux et les classes "
     "du composant avec leurs attributs. Les stéréotypes indiquent le rôle de la classe : «entity» (donnée persistante), "
     "«service» (règles métier), «control» (règle technique) et «value object» (objet de valeur).")
for i, p in enumerate(PACKAGES):
    h("4.2.%d Package %s" % (i + 1, p["name"]), 3, brk=(i > 0))
    para("Exigences : %s. %s" % (" · ".join(p["fr"]), p["role"])).paragraph_format.keep_with_next = True
    figure(os.path.join(IMG, "wb_%s_struct.png" % p["name"].lower()), 17.5,
           "White box du package %s : sous-packages, composants et classes" % p["name"], max_h_cm=8.0)
    caption("Composition du package %s" % p["name"], kind="Tableau")
    table(["Sous-package", "Composant", "Classes", "Rôle"], [
        [s["name"], s["comp"], [c["name"] for c in s["classes"]], SUB_ROLE[(p["name"], s["name"])]]
        for s in p["subs"]
    ], [4.2, 5.4, 6.0, 9.8], font_size=10)

# ================================================================== 5
h("5. Étape 3 – Interactions : dépendances et interfaces", brk=True)
para("Les packages ne communiquent que par leurs interfaces : un package n'appelle jamais directement les classes "
     "internes d'un autre. La figure ci-dessous remet les packages en black box et montre, pour chaque dépendance, "
     "l'interface utilisée.")
figure(os.path.join(IMG, "fig02_black_box.png"), 14.5,
       "Dépendances et interfaces entre les packages (black box)", max_h_cm=14.5)
caption("Dépendances entre packages", kind="Tableau")
table(["Source", "Dépend de", "Interfaces utilisées"], [list(d) for d in DEPENDENCIES], [5.5, 7.5, 12.4])
caption("Catalogue des interfaces", kind="Tableau")
table(["Interface", "Fournie par", "Opérations principales", "Utilisée par"],
      [list(i) for i in INTERFACES], [4.0, 4.4, 10.0, 7.0], font_size=9.5, mono_cols=(0,))
h("5.1 Règles d'architecture", 2)
bullets([
    ("Aucun cycle : ", "les dépendances vont toujours vers le bas (AI, Insight, Collaboration → Tasks → Sprints → Projects → Identity). "
                       "Identity ne dépend de rien."),
    ("Inversion par événements : ", "Sprints et Tasks ne connaissent pas Collaboration ; ils publient leurs événements sur "
                                    "IDomainEvents (fourni par Projects), et Collaboration s'y abonne pour alimenter l'historique et les "
                                    "notifications (patron observateur). Cela évite un cycle Tasks ↔ Collaboration."),
    ("Contrôle d'accès centralisé : ", "toute règle « cet utilisateur peut-il voir / modifier ce projet ? » passe par IProjectAccess."),
    ("IA optionnelle : ", "seul le package AI dépend du service IA externe (IAiApi), lui-même dépendant d'un fournisseur LLM "
                          "facultatif ; sans IA, les six autres packages fonctionnent."),
    ("Backend passerelle : ", "le navigateur ne parle qu'au backend ; le service IA n'est jamais exposé et n'accède pas à la base."),
])

# ================================================================== 6
h("6. Étape 4 – Méthodes des classes", brk=True)
para("Chaque classe reçoit maintenant ses méthodes (compartiment du bas). Les noms des services et leurs méthodes reprennent "
     "ceux du backend Express ; les entités reçoivent des méthodes de domaine qui portent leurs règles propres. Les diagrammes "
     "montrent aussi les interfaces fournies (cercle sous le package) et requises (demi-cercle au-dessus).")
for i, p in enumerate(PACKAGES):
    h("6.%d Package %s" % (i + 1, p["name"]), 2, brk=(i > 0))
    figure(os.path.join(IMG, "wb_%s_full.png" % p["name"].lower()), 17.5,
           "Package %s avec les méthodes et les interfaces" % p["name"], max_h_cm=13.0)
    caption("Classes et méthodes du package %s" % p["name"], kind="Tableau").paragraph_format.page_break_before = True
    rows = []
    for s in p["subs"]:
        for c in s["classes"]:
            rows.append([c["name"], c["stereo"], ["+ " + m for m in c.get("methods", [])], CLASS_ROLE[c["name"]]])
    table(["Classe", "Type", "Méthodes", "Responsabilité"], rows, [4.8, 2.6, 8.2, 9.8], font_size=9.5, mono_cols=(2,))

# ================================================================== 7 / 8 (portrait)
new_section(False)
h("7. Traçabilité des exigences")
para("Chaque exigence de l'atelier 1 est portée par au moins un package, un composant et une classe : aucune exigence n'est orpheline.")
caption("Traçabilité exigences → packages → composants → classes", kind="Tableau")
table(["Exigence", "Package", "Sous-packages", "Classes principales"], [list(r) for r in TRACE],
      [1.8, 2.6, 4.6, 7.6], font_size=9.5)

h("8. Conclusion")
para("L'atelier a fait évoluer les exigences fonctionnelles vers une architecture logicielle structurée en quatre niveaux. "
     "Le modèle obtenu comprend :")
caption("Bilan du modèle", kind="Tableau")
table(["Élément", "Nombre"], [
    ["Packages", str(len(PACKAGES))],
    ["Sous-packages (et composants)", str(n_sub)],
    ["Classes", str(n_cls)],
    ["Méthodes", str(n_meth)],
    ["Interfaces (dont 2 d'API externes)", str(len(INTERFACES))],
    ["Dépendances entre packages et systèmes externes", str(len(DEPENDENCIES))],
], [11, 5.6])
table(["Qualité", "Apport de l'architecture"], [
    ["Maintenabilité", "Une modification d'un package n'impacte que les packages qui l'utilisent, par ses interfaces."],
    ["Modularité", "Frontend, backend et service IA restent des blocs indépendants ; l'IA est isolée dans un seul package."],
    ["Sécurité", "Contrôle d'accès centralisé dans Identity (rôles) et Projects (accès au projet)."],
    ["Testabilité", "Services et interfaces explicites : chaque package se teste seul en simulant ses interfaces requises."],
    ["Évolutivité", "Un nouveau module s'ajoute en se branchant sur IDomainEvents et les interfaces existantes."],
], [3.6, 13.0])
para("Ce modèle sert de base à l'atelier suivant (vue building block), qui place ces packages dans les blocs exécutables : "
     "frontend Angular, backend Express.js, service IA FastAPI et base MongoDB.")

doc.save(OUT)
print("saved", os.path.abspath(OUT), "sub", n_sub, "cls", n_cls, "meth", n_meth)
