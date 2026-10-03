"""Test documents built in memory (no binary fixture in the repository)."""

import io

from docx import Document

SPECIFICATION_FR = """# Cahier des charges — Application de gestion de bibliothèque

## 1. Contexte
La bibliothèque municipale souhaite moderniser la gestion de ses prêts. Le système actuel doit être remplacé.

## 2. Fonctionnalités

### 2.1 Gestion des adhérents
- Inscription d'un adhérent (Must)
- Modification du profil de l'adhérent
- Le système doit permettre au bibliothécaire de désactiver un adhérent [Could]

### 2.2 Catalogue
- En tant que lecteur, je veux rechercher un livre par titre ou auteur afin de savoir s'il est disponible (5 pts)
  - La recherche ignore les accents
  - Les résultats sont paginés
- Import du catalogue depuis un fichier CSV (estimation : 3 jours)

### 2.3 Prêts
1. Emprunter un livre
2. Retourner un livre
3. Envoyer un rappel par e-mail avant l'échéance

## 3. Exigences non fonctionnelles
L'application doit être responsive. Les mots de passe doivent être chiffrés avec bcrypt.

## 4. Contraintes techniques
- Déploiement avec Docker sur un serveur Linux
- Rédiger la documentation utilisateur
- Tests unitaires du module de prêts

## 5. Hors périmètre
- Application mobile native

## 6. Backlog
| ID | User story | Priorité | Points |
|----|------------|----------|--------|
| US-10 | En tant qu'administrateur, je veux exporter les statistiques en PDF | Should | 8 |
| US-11 | En tant que lecteur, je veux réserver un livre déjà emprunté | Won't | 3 |
"""

SPECIFICATION_EN = """Product backlog - Recipe sharing website

Accounts:
- As a visitor, I want to create an account so that I can publish recipes [Must]
- As a member, I want to reset my password
Recipes:
- As a member, I want to publish a recipe with photos (8 points)
- As a member, I want to rate a recipe (Could)
Acceptance criteria:
- A rating is between 1 and 5 stars
- A member rates a recipe only once
"""


def make_docx() -> bytes:
    document = Document()
    document.add_heading("Gestion des utilisateurs", level=1)
    document.add_paragraph("Créer un compte utilisateur", style="List Bullet")
    document.add_paragraph("Le système doit envoyer un e-mail de bienvenue.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Fonctionnalité"
    table.cell(0, 1).text = "Priorité"
    table.cell(1, 0).text = "Exporter la liste des utilisateurs"
    table.cell(1, 1).text = "Haute"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_pdf(lines: list[str]) -> bytes:
    """Minimal one-page PDF (Helvetica, ASCII text) with a valid cross-reference table."""

    def escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    content = "BT /F1 12 Tf 72 720 Td 16 TL " + " ".join(f"({escape(line)}) Tj T*" for line in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(output)
    output += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    output += "".join(f"{offset:010d} 00000 n \n" for offset in offsets).encode()
    output += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return output
